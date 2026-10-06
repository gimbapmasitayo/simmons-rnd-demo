# -*- coding: utf-8 -*-
"""발행 기록과 «실적» 화면 — 임원 보고용 한 장.

다섯 질문에 답한다: ① 얼마나 썼나 ② 얼마나 아꼈나 ③ 믿을 수 있나 ④ 품질은 나아졌나 ⑤ 앞으로 얼마나 더 아끼나.

  · 기록은 성적서를 만들 때마다 발행기록.jsonl 에 한 줄(record). 그 시점의 기준 시간(base_min)을 함께 적어
    나중에 기준 시간이 바뀌어도 과거 집계는 당시 값으로 유지된다.
  · 기준 시간은 실측 기록(measurements)의 평균/중앙값으로 자동 계산한다. 실측이 종류별 2건 이상이면 «확정», 아니면 «확인 중»,
    실측이 없어 손으로 넣은 값이면 «추정치». 기준 시간이 바뀌면 history 에 누가·언제·전/후를 남긴다.
  · 같은 성적서번호로 다시 만들면 «수정 재발행» — 사유(입력 오류 / 측정 재시험 / 양식 변경 / 기타). 입력 오류만 품질 이슈로 센다.
  · ?demo=1 은 메모리 안에서만 만든 더미 자료. 파일에 닿지 않는다.
"""
import os, io, re, json, time, datetime, threading, random, statistics

HERE = os.path.dirname(os.path.abspath(__file__))
LOG = os.path.join(HERE, "발행기록.jsonl")
SETTINGS = os.path.join(HERE, "실적설정.json")
ARCHIVE = os.path.join(HERE, "발행")
_lock = threading.Lock()

KIND_NAMES = {"finished": "라돈 완제품", "material": "라돈 원자재", "kc": "매트리스 KC", "ks": "매트리스 KS"}
KIND_ORDER = ("finished", "material", "kc", "ks")
REISSUE_REASONS = ("입력 오류", "측정 재시험", "양식 변경", "기타")
NO_RE = re.compile(r"^\s*([A-Za-z]{1,4})-(\d{2})-(\d{1,4})\s*$")
NOT_REAL = re.compile(r"(test|테스트|임시|연습|sample|샘플|예시|tmp|temp|xxx|aaa)", re.I)

DEFAULT_SETTINGS = {
    "manual": {"finished": 40, "material": 30, "kc": 60, "ks": 90},   # 실측이 없을 때 쓰는 손 입력값(분, 시료 1종) — «추정치»
    "per_sample": {"finished": 10, "material": 8, "kc": 0, "ks": 0},  # 시료 1종 늘 때마다 더하는 분
    "program": 10,                       # 프로그램에 입력하는 시간(분)
    "method": "mean",                    # 실측 기록을 기준 시간으로 바꾸는 방법: mean | median
    "measurements": {k: [] for k in KIND_ORDER},   # [{by, date, minutes, n_samples, memo}]
    "history": [],                       # 기준 시간 변경 이력 [{ts, by, kind, old, new, how}]
    "pre": {"total": 0, "reissue": 0, "from": "", "to": "", "source": ""},   # 도입 전 비교 (관리자 입력)
    "baseline": 0, "baseline_note": "",  # 도입 전 올해 손으로 만든 건수 (비교용)
    "workday_hours": 8,
    "excluded": [],
    "prefix": "L",
}


# ═══════════════════════════════════════════════ 설정
def settings():
    s = json.loads(json.dumps(DEFAULT_SETTINGS))
    if os.path.exists(SETTINGS):
        try:
            got = json.load(io.open(SETTINGS, encoding="utf-8"))
            for k, v in got.items():
                if isinstance(v, dict) and isinstance(s.get(k), dict):
                    s[k].update(v)
                else:
                    s[k] = v
        except Exception:
            pass
    for k in KIND_ORDER:
        s["measurements"].setdefault(k, [])
    return s


def _write_settings(s):
    with _lock:
        tmp = SETTINGS + ".tmp"
        with io.open(tmp, "w", encoding="utf-8") as f:
            json.dump(s, f, ensure_ascii=False, indent=1)
        os.replace(tmp, SETTINGS)


def base_minutes(s, kind):
    """종류별 기준 시간(시료 1종, 분)과 상태.
    돌려주는 값: {"min", "status": 확정|확인 중|추정치, "n", "method", "by": [...], "last": 마지막 측정일}"""
    ms = [m for m in s["measurements"].get(kind, []) if m.get("minutes")]
    per = float(s["per_sample"].get(kind, 0))
    if ms:
        norm = [max(1.0, float(m["minutes"]) - per * max(0, int(m.get("n_samples") or 1) - 1)) for m in ms]
        val = statistics.median(norm) if s.get("method") == "median" else statistics.mean(norm)
        return {"min": round(val, 1), "status": "확정" if len(ms) >= 2 else "확인 중", "n": len(ms),
                "method": "중앙값" if s.get("method") == "median" else "평균",
                "by": sorted({m.get("by", "") for m in ms if m.get("by")}),
                "last": max((m.get("date") or "" for m in ms), default="")}
    return {"min": float(s["manual"].get(kind, 0)), "status": "추정치", "n": 0, "method": "손 입력", "by": [], "last": ""}


def manual_minutes(s, kind, n_samples, base=None):
    b = float(base) if base is not None else base_minutes(s, kind)["min"]
    return b + float(s["per_sample"].get(kind, 0)) * max(0, int(n_samples) - 1)


def status_overall(s):
    """뱃지용: 모든 종류가 확정이면 확정, 하나라도 추정치/확인 중이면 확인 중."""
    st = [base_minutes(s, k)["status"] for k in KIND_ORDER]
    return "확정" if all(x == "확정" for x in st) else "확인 중"


def save_settings(patch, by=""):
    """설정을 고친다. 기준 시간이 바뀌는 항목은 전/후를 history 에 남긴다."""
    s = settings()
    before = {k: base_minutes(s, k)["min"] for k in KIND_ORDER}
    for k in ("manual", "per_sample"):
        if k in patch and isinstance(patch[k], dict):
            for kind, v in patch[k].items():
                if kind in KIND_NAMES:
                    s[k][kind] = max(0.0, float(v or 0))
    if "program" in patch:
        s["program"] = max(0.0, float(patch["program"] or 0))
    if patch.get("method") in ("mean", "median"):
        s["method"] = patch["method"]
    if "measurements" in patch and isinstance(patch["measurements"], dict):
        for kind, lst in patch["measurements"].items():
            if kind not in KIND_NAMES or not isinstance(lst, list):
                continue
            clean = []
            for m in lst:
                try:
                    mins = float(m.get("minutes") or 0)
                except (TypeError, ValueError):
                    continue
                if mins <= 0:
                    continue
                clean.append({"by": str(m.get("by") or "").strip()[:40], "date": parse_issue_date(str(m.get("date") or "")) or "",
                              "minutes": mins, "n_samples": max(1, int(float(m.get("n_samples") or 1))),
                              "memo": str(m.get("memo") or "").strip()[:160]})
            s["measurements"][kind] = clean
    if "pre" in patch and isinstance(patch["pre"], dict):
        p = s.setdefault("pre", dict(DEFAULT_SETTINGS["pre"]))
        for k in ("total", "reissue"):
            if k in patch["pre"]:
                try:
                    p[k] = max(0, int(float(patch["pre"][k] or 0)))
                except (TypeError, ValueError):
                    pass
        for k in ("from", "to"):
            if k in patch["pre"]:
                p[k] = parse_issue_date(str(patch["pre"][k] or "")) or ""
        if "source" in patch["pre"]:
            p["source"] = str(patch["pre"]["source"] or "").strip()[:200]
    if "baseline" in patch:
        s["baseline"] = max(0, int(float(patch["baseline"] or 0)))
    if "baseline_note" in patch:
        s["baseline_note"] = str(patch["baseline_note"] or "")[:120]
    if "workday_hours" in patch:
        s["workday_hours"] = max(1.0, float(patch["workday_hours"] or 8))
    if "prefix" in patch:
        p = re.sub(r"[^A-Za-z]", "", str(patch["prefix"] or ""))[:4].upper()
        s["prefix"] = p or "L"
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    for k in KIND_ORDER:
        after = base_minutes(s, k)
        if abs(after["min"] - before[k]) >= 0.05:
            s.setdefault("history", []).append({"ts": now, "by": (by or "").strip()[:40] or "미입력", "kind": k,
                                                "old": before[k], "new": after["min"], "how": after["method"], "status": after["status"]})
    s["history"] = s.get("history", [])[-200:]
    _write_settings(s)
    return s


def set_excluded(rec_id, flag):
    s = settings()
    ex = set(int(x) for x in s.get("excluded", []))
    (ex.add if flag else ex.discard)(int(rec_id))
    s["excluded"] = sorted(ex)
    _write_settings(s)
    return s


# ═══════════════════════════════════════════════ 기록
def parse_issue_date(s):
    """«2026년 03월 06일» · 2026-03-06 · 2026.3.6 · 20260306 → ISO. 못 읽으면 None."""
    m = re.search(r"(\d{4})\D{0,3}(\d{1,2})\D{0,3}(\d{1,2})", s or "")
    if not m:
        return None
    try:
        return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
    except ValueError:
        return None


def record(kind, header, n_samples=1, n_runs=1, n_photos=0, warnings=0,
           elapsed=0.0, practice=False, ip="", issuer="", host="", src_files=(), out_file="",
           version=1, backdated=False, issued_on=None, warn_items=(), acked=False, reissue_reason=""):
    """발행 한 건. base_min 에 그 시점의 기준 시간(시료 1종)을 박아 둔다 — 과거 집계는 당시 값으로."""
    now = datetime.datetime.now()
    s = settings()
    row = {"id": int(time.time() * 1000), "ts": now.isoformat(timespec="seconds"),
           "kind": kind, "report_no": (header.get("report_no") or "").strip(),
           "sample_title": (header.get("sample_title") or "").strip()[:60],
           "n_samples": int(n_samples), "n_runs": int(n_runs), "n_photos": int(n_photos),
           "warnings": int(warnings), "elapsed": round(float(elapsed), 2),
           "practice": bool(practice), "ip": ip or "",
           "issuer": (issuer or "").strip()[:40], "host": (host or "")[:60],
           "src_files": [str(x)[:80] for x in (src_files or [])][:12], "out_file": out_file or "",
           "version": int(version), "backdated": bool(backdated),
           "issued_on": issued_on or now.date().isoformat(),
           "warn_items": list(warn_items or []), "acked": bool(acked),
           "reissue": int(version) > 1,
           "reissue_reason": (reissue_reason if reissue_reason in REISSUE_REASONS else ("기타" if int(version) > 1 else "")),
           "base_min": base_minutes(s, kind)["min"], "per_min": float(s["per_sample"].get(kind, 0)), "program_min": float(s["program"])}
    with _lock:
        with io.open(LOG, "a", encoding="utf-8") as f:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
    return row


def load():
    if not os.path.exists(LOG):
        return []
    out = []
    with io.open(LOG, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    out.append(json.loads(line))
                except ValueError:
                    pass
    return out


def migrate():
    """옛 기록에 새 구조(base_min · reissue · reissue_reason · issued_on)를 채운다. 바꾸기 전 백업을 남긴다."""
    rows = load()
    if not rows or all("base_min" in r and "reissue" in r for r in rows):
        return 0
    s = settings()
    bak = os.path.join(HERE, "_backup", "발행기록_migrate_%s.jsonl" % datetime.datetime.now().strftime("%Y%m%d_%H%M%S"))
    os.makedirs(os.path.dirname(bak), exist_ok=True)
    import shutil
    shutil.copy(LOG, bak)
    seen = {}
    n = 0
    for r in sorted(rows, key=lambda r: r["ts"]):
        changed = False
        if "base_min" not in r:
            r["base_min"], r["per_min"], r["program_min"] = base_minutes(s, r["kind"])["min"], float(s["per_sample"].get(r["kind"], 0)), float(s["program"])
            changed = True
        if "version" not in r:
            r["version"] = seen.get(r.get("report_no"), 0) + 1
            changed = True
        seen[r.get("report_no")] = max(seen.get(r.get("report_no"), 0), r["version"])
        if "reissue" not in r:
            r["reissue"] = r["version"] > 1
            r["reissue_reason"] = "기타" if r["reissue"] else ""
            changed = True
        r.setdefault("issued_on", r["ts"][:10])
        r.setdefault("warn_items", []); r.setdefault("acked", False)
        n += changed
    with _lock:
        tmp = LOG + ".tmp"
        with io.open(tmp, "w", encoding="utf-8") as f:
            for r in sorted(rows, key=lambda r: r["ts"]):
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
        os.replace(tmp, LOG)
    return n


def set_reason(rec_id, reason):
    if reason not in REISSUE_REASONS:
        return False
    rows = load()
    ok = False
    for r in rows:
        if int(r["id"]) == int(rec_id) and r.get("reissue"):
            r["reissue_reason"] = reason
            ok = True
    if ok:
        with _lock:
            tmp = LOG + ".tmp"
            with io.open(tmp, "w", encoding="utf-8") as f:
                for r in rows:
                    f.write(json.dumps(r, ensure_ascii=False) + "\n")
            os.replace(tmp, LOG)
    return ok


def same_no(report_no):
    no = (report_no or "").strip()
    if not no:
        return []
    return sorted((r for r in load() if (r.get("report_no") or "").strip().lower() == no.lower()), key=lambda r: r["ts"])


def next_number(kind=None, today=None):
    s = settings()
    prefix = (s.get("prefix") or "L").strip() or "L"
    today = today or datetime.date.today()
    yy = today.strftime("%y")
    mx = 0
    for r in load():
        m = NO_RE.match(r.get("report_no") or "")
        if m and m.group(1).lower() == prefix.lower() and m.group(2) == yy:
            mx = max(mx, int(m.group(3)))
    return "%s-%s-%03d" % (prefix, yy, mx + 1), prefix


def archive_path(report_no, version):
    safe = re.sub(r'[\\/:*?"<>|]+', "_", (report_no or "").strip()) or "번호없음"
    d = os.path.join(ARCHIVE, datetime.date.today().strftime("%Y"), safe)
    os.makedirs(d, exist_ok=True)
    return os.path.join(d, "v%d_%s.xlsx" % (version, datetime.datetime.now().strftime("%Y%m%d_%H%M%S")))


# ═══════════════════════════════════════════════ 더미 (?demo=1) — 파일에 닿지 않는다
def demo_rows(year=None):
    year = year or datetime.date.today().year
    rnd = random.Random(2026)
    rows, no = [], 0
    names = ["김연구", "박품질", "이측정", "최시험", ""]
    titles = {"finished": ["특판 5종", "Black 7종", "Edition 11종", "Online 10종", "N32 5종"],
              "material": ["원단 27종", "폼 12종", "부직포 8종"], "kc": ["샘플 1번", "하이브리드 Q"], "ks": ["샘플 2번", "뷰티레스트 K"]}
    warn_pool = ["시료 사진이 없습니다", "시험환경 온도 33℃", "RAD7 42 와 RD200 61 라돈값이 31% 차이 납니다", "의뢰일자가 발급일자보다 늦습니다"]
    base = {"finished": 42, "material": 31, "kc": 58, "ks": 88}
    for mo in range(1, 11):
        for i in range(6):
            kind = KIND_ORDER[(mo + i) % 4] if i < 4 else rnd.choice(KIND_ORDER)
            no += 1
            day = min(28, 2 + i * 4 + rnd.randint(0, 2))
            d = datetime.date(year, mo, day)
            ns = rnd.randint(3, 11) if kind in ("finished", "material") else 1
            warn = rnd.random() < 0.12
            rows.append({"id": 1700000000000 + no, "ts": d.isoformat() + "T10:%02d:00" % rnd.randint(0, 59), "kind": kind,
                         "report_no": "L-%s-%03d" % (str(year)[2:], no), "sample_title": rnd.choice(titles[kind]),
                         "n_samples": ns, "n_runs": 1, "n_photos": 0 if warn else rnd.randint(1, 3), "warnings": int(warn),
                         "elapsed": rnd.uniform(4, 12), "practice": False, "ip": "10.0.0.%d" % (11 + (no % 5)),
                         "issuer": names[no % 5], "host": "PC-%02d" % (1 + no % 5), "src_files": [], "out_file": "",
                         "version": 1, "backdated": False, "issued_on": d.isoformat(),
                         "warn_items": [rnd.choice(warn_pool)] if warn else [], "acked": warn, "reissue": False, "reissue_reason": "",
                         "base_min": base[kind], "per_min": DEFAULT_SETTINGS["per_sample"][kind], "program_min": 10})
    # 재발행 4건
    for j, reason in enumerate(("입력 오류", "입력 오류", "측정 재시험", "양식 변경")):
        src = rows[7 + j * 11]
        r = dict(src, id=src["id"] + 500, ts=src["ts"][:11] + "15:00:00", version=2, reissue=True, reissue_reason=reason, warnings=0, warn_items=[], acked=False)
        rows.append(r)
    rows.sort(key=lambda r: r["ts"])
    s = json.loads(json.dumps(DEFAULT_SETTINGS))
    s["measurements"] = {"finished": [{"by": "김연구", "date": "%d-02-10" % year, "minutes": 85, "n_samples": 5, "memo": "특판 5종 수기 작성"},
                                      {"by": "박품질", "date": "%d-03-04" % year, "minutes": 100, "n_samples": 7, "memo": ""}],
                         "material": [{"by": "이측정", "date": "%d-03-12" % year, "minutes": 86, "n_samples": 8, "memo": "원단 묶음"},
                                      {"by": "김연구", "date": "%d-04-02" % year, "minutes": 38, "n_samples": 2, "memo": ""}],
                         "kc": [{"by": "최시험", "date": "%d-05-20" % year, "minutes": 58, "n_samples": 1, "memo": ""},
                                {"by": "최시험", "date": "%d-06-11" % year, "minutes": 61, "n_samples": 1, "memo": ""}],
                         "ks": [{"by": "최시험", "date": "%d-06-25" % year, "minutes": 92, "n_samples": 1, "memo": "사진 22장 포함"}]}
    s["pre"] = {"total": 48, "reissue": 9, "from": "%d-01-01" % (year - 1), "to": "%d-12-31" % (year - 1), "source": "%d년 성적서 대장 (품질팀)" % (year - 1)}
    s["history"] = [{"ts": "%d-03-05 09:12" % year, "by": "박품질", "kind": "finished", "old": 40, "new": 42.5, "how": "평균", "status": "확정"}]
    return rows, s


# ═══════════════════════════════════════════════ 집계
def is_real(row):
    no = row.get("report_no") or ""
    return bool(no) and not NOT_REAL.search(no) and not row.get("practice")


def bucket_key(d, group):
    if group == "day":
        return d.isoformat(), d.strftime("%m.%d")
    if group == "week":
        y, w, _ = d.isocalendar()
        monday = d - datetime.timedelta(days=d.weekday())
        return "%04d-W%02d" % (y, w), monday.strftime("%m.%d") + "~"
    if group == "month":
        return d.strftime("%Y-%m"), d.strftime("%y.%m")
    if group == "quarter":
        q = (d.month - 1) // 3 + 1
        return "%04d-Q%d" % (d.year, q), "%d년 %dQ" % (d.year, q)
    return "%04d" % d.year, "%d년" % d.year


def _rec_minutes(s, r):
    """기록 한 건의 손 작성 시간·프로그램 시간 — 기록에 박힌 당시 기준값을 우선한다."""
    base = r.get("base_min")
    per = r.get("per_min", s["per_sample"].get(r["kind"], 0))
    prog = r.get("program_min", s["program"])
    if base is None:
        base = base_minutes(s, r["kind"])["min"]
    manual = float(base) + float(per) * max(0, int(r.get("n_samples", 1)) - 1)
    return manual, float(prog)


def _totals(s, inrange):
    """기간 안 기록의 합계. 번호별 한 건(마지막 판이 대표)."""
    by_no = {}
    for r in inrange:
        if r["real"]:
            by_no.setdefault(r["report_no"], []).append(r)
    counted = []
    for no, lst in by_no.items():
        rep = dict(lst[-1])
        rep["regen"] = len(lst) - 1
        rep["manual"], rep["prog"] = _rec_minutes(s, rep)
        rep["saved"] = max(0.0, rep["manual"] - rep["prog"])
        counted.append(rep)
    counted.sort(key=lambda r: r["ts"])
    real = [r for r in inrange if r["real"]]
    reissues = [r for r in real if r.get("reissue")]
    warned = [r for r in real if r.get("warnings") or r.get("warn_items")]
    manual_h = sum(r["manual"] for r in counted) / 60.0
    prog_h = sum(r["prog"] for r in counted) / 60.0
    users = {}
    for r in counted:
        users[(r.get("issuer") or "").strip() or "미입력"] = users.get((r.get("issuer") or "").strip() or "미입력", 0) + 1
    by_reason = {k: 0 for k in REISSUE_REASONS}
    for r in reissues:
        by_reason[r.get("reissue_reason") or "기타"] = by_reason.get(r.get("reissue_reason") or "기타", 0) + 1
    warn_types = {}
    for r in warned:
        for w in (r.get("warn_items") or ["(종류 미기록)"]):
            key = re.split(r"[\d(]", str(w))[0].strip()[:18] or "기타"
            warn_types[key] = warn_types.get(key, 0) + 1
    kinds = {}
    for k in KIND_ORDER:
        lst = [r for r in counted if r["kind"] == k]
        kinds[k] = {"n": len(lst), "avg_manual": round(sum(r["manual"] for r in lst) / len(lst)) if lst else None,
                    "saved_h": round(sum(r["saved"] for r in lst) / 60.0, 1), "manual_h": round(sum(r["manual"] for r in lst) / 60.0, 1)}
    return {"count": len(counted), "by_kind": kinds,
            "saved_h": round(manual_h - prog_h, 1), "manual_h": round(manual_h, 1), "program_h": round(prog_h, 1),
            "pct": round(100.0 * (1 - prog_h / manual_h)) if manual_h > 0 else 0,
            "gens": len(real), "reissue": len(reissues), "reissue_pct": round(100.0 * len(reissues) / len(counted)) if counted else 0,
            "reissue_err": by_reason.get("입력 오류", 0), "by_reason": by_reason,
            "warned": len(warned), "warned_acked": sum(1 for r in warned if r.get("acked")), "warn_types": warn_types,
            "clean_pct": round(100.0 * sum(1 for r in counted if not (r.get("warnings") or r.get("warn_items"))) / len(counted)) if counted else None,
            "users": sorted(users.items(), key=lambda x: -x[1]), "n_users": len([u for u in users if u != "미입력"]),
            "pcs": len({r.get("host") or r.get("ip") for r in counted if (r.get("host") or r.get("ip"))}),
            "samples": sum(int(r.get("n_samples", 1)) for r in counted),
            "first": counted[0]["issued_on"] if counted else None, "last": counted[-1]["issued_on"] if counted else None}, counted


def summarize(rows, s, d_from, d_to, group="month", demo=False):
    ex = set(int(x) for x in s.get("excluded", []))
    ex_nos = set(r.get("report_no") for r in rows if int(r["id"]) in ex and r.get("report_no"))

    def prep(lo, hi):
        out = []
        for r in rows:
            try:
                d = datetime.date.fromisoformat((r.get("issued_on") or r["ts"])[:10])
            except Exception:
                continue
            if lo <= d <= hi:
                r = dict(r)
                r["date"] = d
                r["excluded"] = int(r["id"]) in ex or (r.get("report_no") in ex_nos)
                r["real"] = is_real(r) and not r["excluded"]
                out.append(r)
        out.sort(key=lambda r: r["ts"])
        return out

    inrange = prep(d_from, d_to)
    total, counted = _totals(s, inrange)
    # 직전 같은 길이 기간 — 증감 표시용
    span = (d_to - d_from).days + 1
    prev_total, _ = _totals(s, prep(d_from - datetime.timedelta(days=span), d_from - datetime.timedelta(days=1)))
    delta = {"count": total["count"] - prev_total["count"], "saved_h": round(total["saved_h"] - prev_total["saved_h"], 1),
             "pct": total["pct"] - prev_total["pct"], "reissue": total["reissue"] - prev_total["reissue"],
             "n_users": total["n_users"] - prev_total["n_users"], "prev_count": prev_total["count"]}

    buckets = {}
    def B(k, label):
        return buckets.setdefault(k, {"key": k, "label": label, "gens": 0, "count": {kk: 0 for kk in KIND_ORDER},
                                      "saved": 0.0, "manual": 0.0, "reissue": 0, "warned": 0})
    for r in inrange:
        k, label = bucket_key(r["date"], group)
        b = B(k, label)
        b["gens"] += 1
        if r["real"] and r.get("reissue"):
            b["reissue"] += 1
        if r["real"] and (r.get("warnings") or r.get("warn_items")):
            b["warned"] += 1
    for r in counted:
        k, label = bucket_key(r["date"], group)
        b = B(k, label)
        b["count"][r["kind"]] += 1
        b["saved"] += r["saved"]
        b["manual"] += r["manual"]
    blist = [buckets[k] for k in sorted(buckets)]
    cum = 0.0
    for b in blist:
        b["total"] = sum(b["count"].values())
        b["saved_h"] = round(b["saved"] / 60.0, 1)
        b["manual_h"] = round(b["manual"] / 60.0, 1)
        cum += b["saved"]
        b["cum_h"] = round(cum / 60.0, 1)

    # 전망: 최근 3개월 평균 절감 × 남은 개월 (올해 기간을 볼 때만, 월 자료 3개월 이상)
    forecast = None
    today = datetime.date.today()
    months = {}
    for r in counted:
        months[r["date"].strftime("%Y-%m")] = months.get(r["date"].strftime("%Y-%m"), 0.0) + r["saved"]
    if d_to.year == today.year and d_from.month == 1 and len(months) >= 3:
        recent = [months.get((today.replace(day=1) - datetime.timedelta(days=30 * i)).strftime("%Y-%m"), 0.0) for i in range(1, 4)]
        per_month = sum(recent) / 3.0 / 60.0
        remain = 12 - today.month
        forecast = {"per_month_h": round(per_month, 1), "remain": remain,
                    "year_h": round(total["saved_h"] + per_month * remain, 1)}

    recent_rows = []
    for r in reversed(inrange[-300:]):
        manual, prog = _rec_minutes(s, r)
        recent_rows.append({"id": r["id"], "ts": r["ts"][:16].replace("T", " "), "kind": r["kind"],
                            "kind_name": KIND_NAMES.get(r["kind"], r["kind"]),
                            "report_no": r["report_no"], "sample_title": r["sample_title"],
                            "n_samples": r["n_samples"], "n_photos": r["n_photos"],
                            "warnings": int(bool(r.get("warnings") or r.get("warn_items"))), "warn_items": r.get("warn_items", []), "acked": bool(r.get("acked")),
                            "practice": bool(r.get("practice")), "excluded": r["excluded"], "real": r["real"],
                            "issued_on": r.get("issued_on", r["ts"][:10]), "backdated": bool(r.get("backdated")),
                            "issuer": r.get("issuer", ""), "host": r.get("host") or r.get("ip", ""),
                            "version": r.get("version", 1), "reissue": bool(r.get("reissue")), "reissue_reason": r.get("reissue_reason", ""),
                            "has_file": bool(r.get("out_file")) and os.path.exists(os.path.join(HERE, r.get("out_file", ""))),
                            "src_files": r.get("src_files", []), "base_min": round(manual), "saved_min": round(max(0, manual - prog)),
                            "why": ("연습용" if r.get("practice") else "번호 없음" if not r["report_no"] else
                                    "테스트 번호" if NOT_REAL.search(r["report_no"]) else "제외함" if r["excluded"] else "")})

    bases = {k: base_minutes(s, k) for k in KIND_ORDER}
    pre = s.get("pre") or {}
    pre_block = None
    if pre.get("total") and pre.get("reissue") is not None:
        pre_block = {"total": pre["total"], "reissue": pre["reissue"], "pct": round(100.0 * pre["reissue"] / pre["total"]),
                     "from": pre.get("from", ""), "to": pre.get("to", ""), "source": pre.get("source", "")}
    try:
        import verify
        v = verify.result()
        vres = {"ok": v.get("ok", 0), "n": v.get("n", 0), "at": v.get("at", ""), "stale": v.get("stale", False)}
    except Exception:
        vres = {"ok": 0, "n": 0, "at": "", "stale": False}
    return {"buckets": blist, "total": total, "delta": delta, "recent": recent_rows, "group": group,
            "from": d_from.isoformat(), "to": d_to.isoformat(), "kinds": [(k, KIND_NAMES[k]) for k in KIND_ORDER],
            "settings": {"program": s["program"], "per_sample": s["per_sample"], "workday_hours": s.get("workday_hours", 8),
                         "baseline": s.get("baseline", 0), "baseline_note": s.get("baseline_note", ""), "method": s.get("method", "mean"),
                         "prefix": s.get("prefix", "L"), "history": s.get("history", [])[-30:][::-1]},
            "bases": bases, "status": status_overall(s), "pre": pre_block, "forecast": forecast, "verify": vres,
            "demo": bool(demo), "reasons": list(REISSUE_REASONS)}


def data_for(request_args):
    """라우트에서 쓰는 한 줄. demo=1 이면 더미."""
    from datetime import date as _d
    demo = request_args.get("demo") == "1"
    try:
        d_from = _d.fromisoformat(request_args.get("from", ""))
    except ValueError:
        d_from = _d(_d.today().year, 1, 1)
    try:
        d_to = _d.fromisoformat(request_args.get("to", ""))
    except ValueError:
        d_to = _d.today()
    group = request_args.get("group", "month")
    if group not in ("day", "week", "month", "quarter", "year"):
        group = "month"
    if demo:
        rows, s = demo_rows()
        return summarize(rows, s, d_from, d_to, group, demo=True)
    return summarize(load(), settings(), d_from, d_to, group)


# ═══════════════════════════════════════════════ 엑셀 내려받기
def export_xlsx(rows, s):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "발행 기록"
    head = ["성적서번호", "종류", "발급일", "만든 시각", "발행자", "PC", "시료명", "시료 수", "기준 시간(분)", "프로그램(분)", "절감(분)",
            "경고", "경고 항목", "확인 후 발행", "재발행", "재발행 사유", "판", "소급", "실적 포함", "원본 파일"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    ex = set(int(x) for x in s.get("excluded", []))
    for r in sorted(rows, key=lambda r: (r.get("issued_on") or r["ts"])):
        manual, prog = _rec_minutes(s, r)
        ws.append([r.get("report_no", ""), KIND_NAMES.get(r["kind"], r["kind"]), r.get("issued_on", r["ts"][:10]), r["ts"].replace("T", " "),
                   r.get("issuer", ""), r.get("host") or r.get("ip", ""), r.get("sample_title", ""), r.get("n_samples", 1),
                   round(manual), round(prog), round(max(0, manual - prog)),
                   "있음" if (r.get("warnings") or r.get("warn_items")) else "", " / ".join(r.get("warn_items") or []),
                   "예" if r.get("acked") else "", "예" if r.get("reissue") else "", r.get("reissue_reason", ""), r.get("version", 1),
                   "예" if r.get("backdated") else "", "" if (is_real(r) and int(r["id"]) not in ex) else "아니오", ", ".join(r.get("src_files") or [])])
    widths = [12, 12, 11, 17, 10, 10, 22, 7, 12, 11, 9, 6, 40, 10, 8, 11, 5, 6, 9, 50]
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ═══════════════════════════════════════════════ 화면
PAGE = r"""<!doctype html>
<html lang="ko"><head><meta charset="utf-8"><title>연구소 성적서 자동화 · 실적</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
<link rel="stylesheet" href="https://cdn.jsdelivr.net/gh/orioncactus/pretendard@v1.3.9/dist/web/variable/pretendardvariable-dynamic-subset.min.css" onerror="this.remove()">
<style>
:root{--ink:#0f0f11;--ink-2:#3a3a40;--mute:#84848a;--faint:#b4b4b9;--line:#e3e0d9;--line-2:#eeebe4;--paper:#fff;--bg:#f1efea;--black:#0a0a0b;--ok:#1f5a3a;--ok-soft:#e6efe9;--g1:#0f0f11;--g2:#55555b;--g3:#9a9a9f;--g4:#cfcfd3}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);font:14px/1.55 "Pretendard Variable",Pretendard,"Malgun Gothic",system-ui,sans-serif}
__NAVCSS__
.wrap{max-width:1150px;margin:0 auto;padding:36px 36px 60px}
.head{display:flex;justify-content:space-between;align-items:flex-end;gap:16px;flex-wrap:wrap}
.eyebrow{font-size:11px;letter-spacing:.2em;color:var(--mute);font-weight:600}
h1{font-size:26px;font-weight:700;letter-spacing:-.02em;margin:3px 0 0}
.tools{display:flex;gap:6px}
.tools button,.tools a{font:inherit;font-size:12px;padding:7px 12px;border:1px solid var(--line);background:#fff;color:var(--ink);cursor:pointer;text-decoration:none}
.tools button:hover,.tools a:hover{border-color:var(--black)}
.lead{font-size:13px;color:var(--ink-2);margin:8px 0 12px;max-width:820px;line-height:1.6}
.filters{display:flex;flex-wrap:wrap;gap:6px 14px;align-items:center;font-size:12px;color:var(--mute);margin-bottom:16px}
.seg{display:inline-flex;border:1px solid var(--line);background:#fff}
.seg button{font:inherit;font-size:11.5px;padding:4px 9px;background:transparent;border:none;border-right:1px solid var(--line);cursor:pointer;color:var(--ink-2)}
.seg button:last-child{border-right:none}.seg button.on{background:var(--black);color:#fff}
.filters input[type=date]{font:inherit;font-size:11.5px;padding:3px 6px;border:1px solid var(--line);background:#fff}
.demo{background:var(--ink);color:#fff;font-size:12px;padding:6px 14px;margin:-10px 0 14px;letter-spacing:.04em}
.demo a{color:#fff;margin-left:10px}
/* KPI */
.kpis{display:grid;grid-template-columns:repeat(5,1fr);gap:1px;background:var(--line);border:1px solid var(--line)}
.kpi{background:var(--paper);padding:16px 18px 14px;position:relative}
.kpi .k{font-size:11.5px;color:var(--mute);font-weight:600}
.kpi .v{font-size:44px;font-weight:700;letter-spacing:-.035em;line-height:1.05;margin:6px 0 4px;font-variant-numeric:tabular-nums;color:var(--ink)}
.kpi .v small{font-size:14px;font-weight:500;color:var(--mute);margin-left:3px;letter-spacing:0}
.kpi .v.none{font-size:30px;font-weight:400;color:var(--faint)}
.kpi .d{font-size:11.5px;color:var(--mute)}
.kpi .d b{color:var(--ink-2);font-weight:600}
.kpi[data-tip]:hover::after{content:attr(data-tip);position:absolute;left:14px;bottom:-30px;background:var(--black);color:#fff;font-size:11.5px;padding:5px 9px;white-space:nowrap;z-index:5}
/* 뱃지 */
.badges{display:flex;flex-wrap:wrap;gap:8px;margin:12px 0 18px}
.badge{font-size:12px;padding:6px 12px;border:1px solid var(--line);background:#fff;cursor:pointer;color:var(--ink-2);display:inline-flex;gap:6px;align-items:center}
.badge i{font-style:normal;width:8px;height:8px;border-radius:50%;background:var(--faint);display:inline-block}
.badge.ok i{background:var(--ok)}.badge.ok{color:var(--ok);border-color:#c9d9cf}
.badge b{font-weight:600}
.bdetail{background:#fff;border:1px solid var(--line);padding:14px 18px;font-size:12.5px;color:var(--ink-2);margin:-10px 0 18px;line-height:1.6;display:none}
.bdetail.open{display:block}
.bdetail table{font-size:12px;width:auto}
/* 카드 */
.card{background:var(--paper);border:1px solid var(--line);padding:18px 20px;margin-bottom:14px}
.card h2{font-size:14px;font-weight:700;margin:0 0 4px}
.card .hint{font-size:12px;color:var(--mute);margin-bottom:12px;line-height:1.5}
/* 비교 막대 */
.cmp{display:grid;grid-template-columns:120px 1fr 90px;gap:8px 14px;align-items:center}
.cmp .lab{font-size:12.5px;color:var(--ink-2);font-weight:600}
.cmp .bar{height:30px;background:var(--line-2);display:flex;overflow:hidden}
.cmp .bar i{display:block;height:100%}
.cmp .num{font-size:15px;font-weight:700;text-align:right;font-variant-numeric:tabular-nums}
.cmp .num small{font-size:11px;color:var(--mute);font-weight:500;margin-left:2px}
.arrow{grid-column:1/-1;text-align:center;font-size:15px;font-weight:700;color:var(--ok);padding:2px 0}
.legend{display:flex;gap:14px;flex-wrap:wrap;font-size:11.5px;color:var(--mute);margin-top:10px}
.legend i{display:inline-block;width:10px;height:10px;margin-right:5px;vertical-align:-1px}
/* 차트 블록 */
.two{display:grid;grid-template-columns:2fr 1fr;gap:14px}
.tabs{display:flex;gap:2px;border-bottom:1px solid var(--line);margin-bottom:10px}
.tabs button{font:inherit;font-size:12.5px;font-weight:600;padding:8px 12px;background:none;border:none;border-bottom:2px solid transparent;margin-bottom:-1px;cursor:pointer;color:var(--mute)}
.tabs button.on{color:var(--ink);border-bottom-color:var(--ink)}
svg{width:100%;height:auto;display:block}
.gl{stroke:var(--line-2);stroke-width:1}.ax{font-size:10.5px;fill:var(--mute)}.lbl{font-size:11px;font-weight:600;fill:var(--ink)}.lbl.ok{fill:var(--ok)}
.fc{font-size:12.5px;color:var(--ink-2);margin-top:6px}
.fc b{color:var(--ok)}
.klist{display:flex;flex-direction:column}
.krow{display:grid;grid-template-columns:1fr auto;gap:0 10px;padding:9px 0;border-bottom:1px solid var(--line-2);align-items:baseline}
.krow:last-child{border-bottom:none}
.krow .n{font-size:13px;font-weight:600}
.krow .n i{display:inline-block;width:9px;height:9px;margin-right:7px;vertical-align:0}
.krow .c{font-size:20px;font-weight:700;font-variant-numeric:tabular-nums}
.krow .c small{font-size:11px;color:var(--mute);font-weight:500;margin-left:2px}
.krow .s{grid-column:1/-1;font-size:11.5px;color:var(--mute)}
.users{margin-top:14px}
.urow{display:grid;grid-template-columns:70px 1fr 36px;gap:0 10px;align-items:center;font-size:12px;padding:3px 0}
.urow .ub{height:10px;background:var(--line-2)}.urow .ub i{display:block;height:100%;background:var(--g2)}
.urow .un{text-align:right;font-variant-numeric:tabular-nums;color:var(--ink-2)}
/* 하단 탭 */
table{width:100%;border-collapse:collapse;font-size:12.5px}
th{font-size:10.5px;letter-spacing:.1em;color:var(--mute);font-weight:600;text-align:left;padding:7px 8px;border-bottom:1px solid var(--line);white-space:nowrap}
td{padding:7px 8px;border-bottom:1px solid var(--line-2);vertical-align:top}
td.n,th.n{text-align:right;font-variant-numeric:tabular-nums}
tr.dim td{color:var(--faint)}
.why{display:inline-block;font-size:10.5px;padding:0 6px;border-radius:9px;background:var(--line-2);color:var(--mute);margin-left:6px}
.why.err{background:#eeebe4;color:var(--ink-2)}
.why.ok{background:var(--ok-soft);color:var(--ok)}
button.x{font:inherit;font-size:11px;padding:3px 8px;border:1px solid var(--line);background:#fff;cursor:pointer;color:var(--ink-2)}
select.rs{font:inherit;font-size:11.5px;padding:2px 4px;border:1px solid var(--line);background:#fff}
.empty{padding:28px;text-align:center;color:var(--mute);line-height:1.7}
.guide{background:var(--paper);border:1px solid var(--line);padding:34px;text-align:center;color:var(--ink-2);line-height:1.8;margin-bottom:14px}
.guide a{color:var(--ink);font-weight:600}
.qsum{font-size:14px;font-weight:600;margin-bottom:12px}
.qsum b{color:var(--ok)}
.qgrid{display:grid;grid-template-columns:1fr 1fr 1fr;gap:14px}
.qbox h3{font-size:12px;color:var(--mute);font-weight:600;margin:0 0 6px;letter-spacing:.06em}
.rules{font-size:12.5px;color:var(--ink-2);line-height:1.75}
.foot{font-size:12px;color:var(--mute);margin-top:14px}
.tip{position:fixed;display:none;background:var(--black);color:#fff;font-size:12px;padding:6px 10px;pointer-events:none;z-index:9}
.pfoot{display:none}
@media(max-width:1250px){.kpi .v{font-size:36px}.kpi .d{font-size:11px}}
@media(max-width:1000px){.kpis{grid-template-columns:repeat(3,1fr)}.two{grid-template-columns:1fr}.qgrid{grid-template-columns:1fr}.wrap{padding:24px 16px 50px}}
/* 발표 모드 */
body.present .side,body.present .filters,body.present .tools,body.present .bottom,body.present .lead,body.present .foot,body.present .bdetail,body.present .two .right,body.present .tabs{display:none!important}
body.present{padding-left:0}
body.present .wrap{max-width:1400px;padding:30px 50px}
body.present .kpi .v{font-size:56px}
body.present .two{grid-template-columns:1fr}
body.present .pfoot{display:block;font-size:12px;color:var(--mute);margin-top:10px}
@media print{@page{size:A4 portrait;margin:12mm}
  .side,.filters,.tools,.bottom,.lead,.bdetail,.two .right,.tabs,.tip,.demo{display:none!important}
  body{background:#fff;padding-left:0;font-size:12px}.wrap{padding:0;max-width:none}
  .kpis{grid-template-columns:repeat(5,1fr)}.kpi{padding:10px 12px}.kpi .v{font-size:26px}.kpi .k,.kpi .d{font-size:10px}.card{break-inside:avoid;padding:12px 14px;margin-bottom:10px}.two{grid-template-columns:1fr}
  .pfoot{display:block;font-size:10.5px;color:var(--mute);margin-top:8px;border-top:1px solid var(--line);padding-top:6px}}
</style></head><body>
__NAV__
<div class="wrap">
<div class="demo" id="demoBar" hidden>데모 데이터 — 실제 집계가 아닙니다 <a href="/실적">실제 자료로</a></div>
<div class="head">
  <div><div class="eyebrow">시몬스 연구소</div><h1 id="title">연구소 성적서 자동화</h1></div>
  <div class="tools"><button id="present">발표 모드</button><button id="print">인쇄 · PDF</button><a id="xlsx" href="/실적/엑셀">엑셀 내려받기</a></div>
</div>
<div class="lead">측정 파일을 올리면 계산·판정·양식 작성이 자동으로 끝나는 성적서 프로그램의 실적입니다. 아래 숫자는 서버에 남은 발행 기록으로 집계했습니다.</div>
<div class="filters" id="filters">
  <span class="seg" id="quick"><button data-q="month">이번 달</button><button data-q="90">최근 3개월</button><button data-q="year" class="on">올해</button><button data-q="all">도입 이후 전체</button></span>
  <span>직접 <input type="date" id="from"> ~ <input type="date" id="to"></span>
  <span>묶음 <span class="seg" id="group"><button data-g="week">주</button><button data-g="month" class="on">월</button><button data-g="quarter">분기</button></span></span>
</div>

<div class="kpis" id="kpis"></div>
<div class="badges" id="badges"></div>
<div class="bdetail" id="bdetail"></div>

<div id="guide" class="guide" hidden></div>
<div id="body">
<div class="card"><h2>손으로 만들 때와 비교</h2><div class="hint">이 기간에 발행한 성적서를 모두 손으로 만들었다면 걸렸을 시간과, 프로그램으로 실제 들인 시간입니다.</div>
  <div class="cmp" id="cmp"></div><div class="legend" id="legend"></div></div>

<div class="two">
  <div class="card left"><div class="tabs" id="ctabs"><button data-c="cum" class="on">월별 절감 시간 누적</button><button data-c="cnt">월별 발행 건수</button><button data-c="share">종류별 비중</button></div>
    <div id="chart"></div><div class="fc" id="fc"></div></div>
  <div class="card right"><h2>종류별</h2><div class="klist" id="klist"></div>
    <div class="users"><h2 style="font-size:12.5px;margin:12px 0 6px">사용 현황</h2><div id="users"></div></div></div>
</div>
<div class="card" id="qline" style="padding:12px 20px"><div class="qsum" id="qsum" style="margin:0"></div></div>
</div>

<div class="bottom">
<div class="card"><div class="tabs" id="btabs"><button data-b="rec" class="on">발행 기록</button><button data-b="tbl">기간별 표</button><button data-b="qual">품질 기록</button><button data-b="rule">집계 기준</button></div>
  <div id="bview"></div></div>
</div>
<div class="pfoot" id="pfoot"></div>
<div class="foot">기록은 성적서를 만들 때마다 서버에 자동으로 남습니다. 소급 등록한 성적서는 발급일자 기준으로 집계합니다.</div>
</div>
<div class="tip" id="tip"></div>
<script>
const $=(s,r=document)=>r.querySelector(s), $$=(s,r=document)=>[...r.querySelectorAll(s)];
const DEMO=new URLSearchParams(location.search).get('demo')==='1';
let G='month', D=null, CT='cum', BT='rec', OPENB='';
const GC=['var(--g1)','var(--g2)','var(--g3)','var(--g4)'];
function iso(d){return d.getFullYear()+'-'+String(d.getMonth()+1).padStart(2,'0')+'-'+String(d.getDate()).padStart(2,'0');}
function esc(s){return String(s??'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));}
const kor=d=>{const [y,m]=d.split('-'); return `${y}년 ${+m}월`;};
const sign=n=>n>0?`▲${n}`:n<0?`▼${Math.abs(n)}`:'변동 없음';
function setQuick(q){
  const t=new Date(); let f;
  if(q==='month') f=new Date(t.getFullYear(),t.getMonth(),1);
  else if(q==='90'){ f=new Date(t); f.setMonth(f.getMonth()-3); }
  else if(q==='year') f=new Date(t.getFullYear(),0,1);
  else f=new Date(2026,0,1);
  $('#from').value=iso(f); $('#to').value=iso(t);
  $$('#quick button').forEach(b=>b.classList.toggle('on',b.dataset.q===q));
  G = q==='month'?'week' : q==='all'?'quarter' : 'month';
  $$('#group button').forEach(x=>x.classList.toggle('on',x.dataset.g===G));
}
async function load(){
  const r=await fetch(`/실적/data?from=${$('#from').value}&to=${$('#to').value}&group=${G}${DEMO?'&demo=1':''}`);
  D=await r.json(); render();
}
function render(){
  const T=D.total, S=D.settings, dl=D.delta, empty=!T.count;
  $('#demoBar').hidden=!D.demo; $('#xlsx').href='/실적/엑셀'+(D.demo?'?demo=1':'');
  const sameYear=D.from.slice(0,4)===D.to.slice(0,4);
  $('#title').textContent='연구소 성적서 자동화 · '+(sameYear?D.from.slice(0,4):D.from.slice(0,4)+'~'+D.to.slice(0,4));
  const wd=S.workday_hours||8;
  const K=(k,v,u,d,tip)=>`<div class="kpi" ${tip?`data-tip="${esc(tip)}"`:''}><div class="k">${k}</div><div class="v ${v==='–'?'none':''}">${v}<small>${u}</small></div><div class="d">${d}</div></div>`;
  const dlt=(v,u)=>dl.prev_count?`직전 같은 기간 대비 <b>${sign(v)}${u}</b>`:'직전 기간 기록 없음';
  $('#kpis').innerHTML= empty ? [['자동 작성 성적서','건'],['절감 시간','시간'],['절감률','%'],['수정 재발행','건'],['사용 중인 담당자','명']].map(([k,u])=>K(k,'–',u,'기록 없음')).join('') :
    K('자동 작성 성적서',T.count.toLocaleString(),'건',dlt(dl.count,'건'))+
    K('절감 시간',T.saved_h.toLocaleString(),'시간',dlt(dl.saved_h,'시간'),`하루 ${wd}시간 기준 약 ${(T.saved_h/wd).toFixed(T.saved_h/wd>=10?0:1)}일`)+
    K('절감률',T.pct,'%',`손 작성 ${T.manual_h.toLocaleString()}시간 → ${T.program_h.toLocaleString()}시간`)+
    K('수정 재발행',T.reissue,'건',`${T.count?`전체의 ${T.reissue_pct}%`:''} · 입력 오류 <b>${T.reissue_err}건</b>`)+
    K('사용 중인 담당자',T.n_users,'명',`PC <b>${T.pcs}대</b>${T.users.find(u=>u[0]==='미입력')?` · 이름 미입력 ${T.users.find(u=>u[0]==='미입력')[1]}건`:''}`);
  // 뱃지
  const V=D.verify, confirmed=D.status==='확정';
  $('#badges').innerHTML=`<span class="badge ${V.n&&V.ok===V.n&&!V.stale?'ok':''}" data-b="verify"><i></i>계산 검증 <b>${V.n?`${V.ok}/${V.n} 일치`:'아직 점검 전'}</b>${V.at?` · ${esc(V.at.slice(0,10))}`:''}${V.stale?' · 코드 변경 후 재점검 필요':''}</span>
    <span class="badge ${confirmed?'ok':''}" data-b="base"><i></i>${confirmed?'기준 시간 확정':'기준 시간 확인 중'}${confirmed?` · <b>${esc([...new Set(Object.values(D.bases).flatMap(b=>b.by))].join(', ')||'실측 기록')}</b> · ${esc(Object.values(D.bases).map(b=>b.last).filter(Boolean).sort().pop()||'')}`:''}</span>
    ${T.clean_pct!==null?`<span class="badge ${T.clean_pct>=90?'ok':''}" data-b="clean"><i></i>경고 없이 완성 <b>${T.clean_pct}%</b></span>`:''}`;
  renderBadgeDetail();
  $('#guide').hidden=!empty; $('#body').style.display=empty?'none':'';
  if(empty){ $('#guide').innerHTML=`아직 발행 기록이 없습니다.<br>올해 발행한 성적서는 성적서 자동화 화면의 <a href="/성적서">「소급 등록」</a>으로 채울 수 있습니다 →`; $('#pfoot').textContent=pfoot(); renderBottom(); return; }
  // 비교 막대 (종류별 농도 구간)
  const mx=Math.max(T.manual_h,T.program_h,0.1);
  const segs=(key)=>D.kinds.map(([k],i)=>{const h=key==='manual'?T.by_kind[k].manual_h:(T.by_kind[k].n*S.program/60); return `<i style="width:${100*h/mx}%;background:${GC[i]}" title="${D.kinds[i][1]} ${h.toFixed(1)}시간"></i>`;}).join('');
  $('#cmp').innerHTML=`<div class="lab">손으로 만들었다면</div><div class="bar">${segs('manual')}</div><div class="num">${T.manual_h.toLocaleString()}<small>시간</small></div>
    <div class="arrow">→ 절감 ${T.saved_h.toLocaleString()}시간 (${T.pct}%)</div>
    <div class="lab">프로그램으로</div><div class="bar">${segs('prog')}</div><div class="num">${T.program_h.toLocaleString()}<small>시간</small></div>`;
  $('#legend').innerHTML=D.kinds.map(([k,n],i)=>`<span><i style="background:${GC[i]}"></i>${n}</span>`).join('');
  renderChart();
  $('#klist').innerHTML=D.kinds.map(([k,n],i)=>{const b=T.by_kind[k]; return `<div class="krow"><span class="n"><i style="background:${GC[i]}"></i>${n}</span><span class="c">${b.n}<small>건</small></span><span class="s">${b.n?`건당 손 작성 평균 ${b.avg_manual}분 → 프로그램 ${S.program}분 · 절감 ${b.saved_h}시간`:'이 기간 발행 없음'}</span></div>`;}).join('');
  const umax=Math.max(1,...T.users.map(u=>u[1]));
  $('#users').innerHTML=T.users.slice(0,5).map(([u,n])=>`<div class="urow"><span title="${esc(u)}" style="overflow:hidden;text-overflow:ellipsis;white-space:nowrap;color:${u==='미입력'?'var(--faint)':'inherit'}">${esc(u)}</span><span class="ub"><i style="width:${100*n/umax}%${u==='미입력'?';background:var(--g4)':''}"></i></span><span class="un">${n}</span></div>`).join('')+`<div style="font-size:11.5px;color:var(--mute);margin-top:6px">담당자 ${T.n_users}명 · PC ${T.pcs}대</div>`;
  $('#qsum').innerHTML=`${T.count}건 중 입력 오류 재발행 <b>${T.reissue_err}건</b>(${T.count?Math.round(100*T.reissue_err/T.count):0}%) · 경고로 사전에 잡은 입력 <b>${T.warned}건</b>${T.warned_acked?` (그중 확인 후 발행 ${T.warned_acked}건)`:''}`;
  $('#pfoot').textContent=pfoot();
  renderBottom();
}
function pfoot(){ const b=D.bases; return `집계 기준: 기준 시간 ${D.kinds.map(([k,n])=>`${n} ${b[k].min}분(${b[k].status})`).join(' · ')} · 프로그램 입력 ${D.settings.program}분 · 생성 ${new Date().toLocaleString('ko-KR')}${D.demo?' · 데모 데이터':''}`; }
function renderBadgeDetail(){
  const el=$('#bdetail'); if(!OPENB){ el.className='bdetail'; return; }
  const V=D.verify, b=D.bases, T=D.total;
  let h='';
  if(OPENB==='verify') h=`과거에 사람이 발행한 성적서 4종(라돈 완제품·원자재, 매트리스 KC·KS)의 값과 프로그램 계산값을 항목별로 대조한 결과입니다. ${V.n?`${V.ok}/${V.n} 일치 (${esc(V.at)})`:'아직 점검하지 않았습니다.'} 상세 표와 재점검은 <a href="/실적/설정#verify">기준 시간 설정 → 계산 검증</a>에 있습니다.`;
  if(OPENB==='base') h=`<table><tr><th>종류</th><th class="n">기준 시간(분)</th><th>상태</th><th>근거</th></tr>${D.kinds.map(([k,n])=>`<tr><td>${n}</td><td class="n">${b[k].min}</td><td>${b[k].status}</td><td>${b[k].n?`실측 ${b[k].n}건 ${b[k].method} · ${esc(b[k].by.join(', '))} · 마지막 ${esc(b[k].last)}`:'실측 기록 없음 — 손으로 넣은 추정치'}</td></tr>`).join('')}</table><div style="margin-top:6px">실측 기록이 종류별 2건 이상이면 «확정»으로 바뀝니다. <a href="/실적/설정">기준 시간 설정</a>에서 입력합니다.</div>`;
  if(OPENB==='clean') h=`이 기간 성적서 ${T.count}건 중 ${T.count-T.warned}건은 입력 이상치 경고 없이 한 번에 완성됐습니다. 경고가 떴던 ${T.warned}건 중 ${T.warned_acked}건은 담당자가 확인하고 그대로 발행했습니다.`;
  el.innerHTML=h; el.className='bdetail open';
}
function ticks(raw){ const mag=Math.pow(10,Math.floor(Math.log10(Math.max(raw,1)/4))), cand=[1,2,2.5,5,10].map(c=>c*mag);
  const step=cand.find(c=>c*4>=raw)||cand[cand.length-1]; return {step, max:step*Math.ceil(Math.max(raw,step)/step)}; }
function frame(W,H,L,R,T,B,max,step){ let s=''; const y=v=>T+(H-T-B)*(1-v/max);
  for(let v=0;v<=max+1e-9;v+=step) s+=`<line class="gl" x1="${L}" x2="${W-R}" y1="${y(v)}" y2="${y(v)}"/><text class="ax" x="${L-6}" y="${y(v)+4}" text-anchor="end">${Number.isInteger(step)?v:v.toFixed(1)}</text>`; return s; }
function renderChart(){
  $$('#ctabs button').forEach(b=>b.classList.toggle('on',b.dataset.c===CT));
  const bk=D.buckets; $('#fc').innerHTML='';
  if(!bk.length){ $('#chart').innerHTML='<div class="empty">이 기간에 기록이 없습니다.</div>'; return; }
  const W=640,H=250,L=36,R=44,T=22,B=28,n=bk.length;
  if(CT==='cum'){
    const F=D.forecast; const lastCum=bk[n-1].cum_h; const top=Math.max(0.1,lastCum,F?F.year_h:0);
    const {step,max}=ticks(top); const extra=F?F.remain:0; const N=n+extra;
    const x=i=>N===1?(L+W-R)/2:L+(W-L-R)*i/(N-1), y=v=>T+(H-T-B)*(1-v/max);
    let s=`<svg viewBox="0 0 ${W} ${H}">`+frame(W,H,L,R,T,B,max,step);
    const pts=bk.map((b,i)=>[x(i),y(b.cum_h)]);
    if(n>1) s+=`<path d="M${pts[0][0]},${y(0)} ${pts.map(p=>`L${p[0]},${p[1]}`).join(' ')} L${pts[n-1][0]},${y(0)} z" fill="var(--ok-soft)"/><path d="${pts.map((p,i)=>(i?'L':'M')+p[0]+','+p[1]).join(' ')}" fill="none" stroke="var(--ok)" stroke-width="2.5" stroke-linejoin="round"/>`;
    if(F&&extra>0){ s+=`<path d="M${pts[n-1][0]},${pts[n-1][1]} L${x(N-1)},${y(F.year_h)}" fill="none" stroke="var(--ok)" stroke-width="2" stroke-dasharray="5 5"/><text class="lbl ok" x="${x(N-1)}" y="${y(F.year_h)-10}" text-anchor="end">전망 ${F.year_h}시간</text>`; }
    const e=pts[n-1]; s+=`<circle cx="${e[0]}" cy="${e[1]}" r="5" fill="var(--ok)" stroke="#fff" stroke-width="2"/><text class="lbl" x="${e[0]}" y="${e[1]-12}" text-anchor="${n>1?'end':'middle'}" dx="${n>1?6:0}">${lastCum}시간</text>`;
    bk.forEach((b,i)=>{ const st=Math.ceil(N/10); if(i%st===0||N<=10) s+=`<text class="ax" x="${pts[i][0]}" y="${H-8}" text-anchor="middle">${esc(b.label)}</text>`; const w=(W-L-R)/Math.max(1,N-1); s+=`<rect x="${pts[i][0]-w/2}" y="${T}" width="${w}" height="${H-T-B}" fill="transparent" data-tip="${esc(b.label)} · 이 기간 ${b.saved_h}시간 · 누적 ${b.cum_h}시간"/>`; });
    $('#chart').innerHTML=s+'</svg>';
    if(F) $('#fc').innerHTML=`이 속도면 올해 약 <b>${F.year_h}시간</b> 절감 예상입니다 (최근 3개월 평균 ${F.per_month_h}시간/월 × 남은 ${F.remain}개월, 점선).`;
  } else if(CT==='cnt'){
    const {step,max}=ticks(Math.max(1,...bk.map(b=>b.total))); const iw=(W-L-R)/n, bw=Math.min(44,iw*0.62);
    const y=v=>T+(H-T-B)*(1-v/max); let s=`<svg viewBox="0 0 ${W} ${H}">`+frame(W,H,L,R,T,B,max,step);
    bk.forEach((b,i)=>{ const x=L+iw*i+(iw-bw)/2; let yy=y(0);
      D.kinds.forEach(([k],ki)=>{ const c=b.count[k]; if(!c) return; const h=y(0)-y(c); s+=`<rect x="${x}" y="${yy-h}" width="${bw}" height="${h}" fill="${GC[ki]}"/>`; yy-=h; });
      if(b.total) s+=`<text class="lbl" x="${x+bw/2}" y="${y(b.total)-6}" text-anchor="middle">${b.total}</text>`;
      const st=Math.ceil(n/10); if(i%st===0||n<=10) s+=`<text class="ax" x="${x+bw/2}" y="${H-8}" text-anchor="middle">${esc(b.label)}</text>`;
      s+=`<rect x="${L+iw*i}" y="${T}" width="${iw}" height="${H-T-B}" fill="transparent" data-tip="${esc(b.label)} · ${b.total}건 (${D.kinds.filter(([k])=>b.count[k]).map(([k,nm])=>nm+' '+b.count[k]).join(', ')||'없음'})"/>`; });
    $('#chart').innerHTML=s+'</svg>'; $('#fc').innerHTML=D.kinds.map(([k,n],i)=>`<span style="margin-right:12px"><i style="display:inline-block;width:9px;height:9px;background:${GC[i]};margin-right:5px"></i>${n}</span>`).join('');
  } else {
    const T=D.total, tot=Math.max(1,T.count); let s=`<svg viewBox="0 0 ${W} ${H}">`; let acc=0; const cx=160, cy=H/2, r=88, ir=56;
    D.kinds.forEach(([k,n],i)=>{ const v=T.by_kind[k].n; if(!v) return; const a0=acc/tot*2*Math.PI-Math.PI/2, a1=(acc+v)/tot*2*Math.PI-Math.PI/2; acc+=v;
      const big=(a1-a0)>Math.PI?1:0; const P=(ang,rr)=>[cx+rr*Math.cos(ang),cy+rr*Math.sin(ang)];
      const [x0,y0]=P(a0,r),[x1,y1]=P(a1,r),[x2,y2]=P(a1,ir),[x3,y3]=P(a0,ir);
      s+=`<path d="M${x0},${y0} A${r},${r} 0 ${big} 1 ${x1},${y1} L${x2},${y2} A${ir},${ir} 0 ${big} 0 ${x3},${y3} z" fill="${GC[i]}" data-tip="${n} ${v}건 (${Math.round(100*v/tot)}%)"/>`; });
    s+=`<text class="lbl" x="${cx}" y="${cy+5}" text-anchor="middle" style="font-size:20px">${T.count}건</text>`;
    D.kinds.forEach(([k,n],i)=>{ const v=T.by_kind[k].n; s+=`<rect x="300" y="${60+i*34}" width="12" height="12" fill="${GC[i]}"/><text class="ax" x="320" y="${71+i*34}" style="font-size:12.5px;fill:var(--ink)">${n}</text><text class="lbl" x="560" y="${71+i*34}" text-anchor="end">${v}건 · ${Math.round(100*v/tot)}%</text>`; });
    $('#chart').innerHTML=s+'</svg>';
  }
}
function renderBottom(){
  $$('#btabs button').forEach(b=>b.classList.toggle('on',b.dataset.b===BT));
  const T=D.total, S=D.settings; let h='';
  if(BT==='rec'){
    const vers={}; D.recent.forEach(r=>{ if(r.report_no) vers[r.report_no]=(vers[r.report_no]||0)+1; });
    h=D.recent.length?`<div style="overflow-x:auto"><table><tr><th>발급일</th><th>종류</th><th>성적서번호</th><th>시료명</th><th>발행자</th><th class="n">시료</th><th class="n">절감(분)</th><th>경고</th><th>재발행</th><th class="n">판</th><th></th></tr>
      ${D.recent.map(r=>`<tr class="${r.real?'':'dim'}"><td>${esc(r.issued_on)}${r.backdated?'<span class="why">소급</span>':''}<div style="font-size:10.5px;color:var(--mute)">${r.ts}</div></td><td>${esc(r.kind_name)}</td>
        <td>${esc(r.report_no)||'—'}${r.why?`<span class="why">${r.why}</span>`:''}</td><td>${esc(r.sample_title)}</td><td>${esc(r.issuer)||'<span style="color:var(--faint)">미입력</span>'}</td>
        <td class="n">${r.n_samples}</td><td class="n">${r.real?r.saved_min:''}</td>
        <td>${r.warnings?`<span class="why err" title="${esc((r.warn_items||[]).join('\n'))}">경고 ${r.warn_items.length||''}${r.acked?' · 확인':''}</span>`:''}</td>
        <td>${r.reissue?`<select class="rs" data-id="${r.id}" ${D.demo?'disabled':''}>${D.reasons.map(x=>`<option ${x===r.reissue_reason?'selected':''}>${x}</option>`).join('')}</select>`:''}</td>
        <td class="n">${r.has_file?`<a href="/실적/파일/${r.id}">v${r.version}</a>`:`v${r.version}`}</td>
        <td>${r.practice||!r.report_no||D.demo?'':`<button class="x" data-id="${r.id}" data-ex="${r.excluded?0:1}">${r.excluded?'포함':'제외'}</button>`}</td></tr>`).join('')}</table></div>
      <div class="hint" style="margin-top:8px">같은 번호로 다시 만들면 재발행으로 기록되고 사유를 여기서 고칩니다. 판(v)을 누르면 그 판 파일을 내려받습니다.</div>`:'<div class="empty">아직 기록이 없습니다.</div>';
  } else if(BT==='tbl'){
    h=D.buckets.length?`<div style="overflow-x:auto"><table><tr><th>기간</th>${D.kinds.map(([k,n])=>`<th class="n">${n}</th>`).join('')}<th class="n">합계</th><th class="n">손 작성(시간)</th><th class="n">절감(시간)</th><th class="n">누적</th><th class="n">재발행</th><th class="n">경고</th></tr>
      ${D.buckets.map(b=>`<tr><td>${esc(b.label)}</td>${D.kinds.map(([k])=>`<td class="n">${b.count[k]||''}</td>`).join('')}<td class="n"><b>${b.total}</b></td><td class="n">${b.manual_h}</td><td class="n">${b.saved_h}</td><td class="n">${b.cum_h}</td><td class="n">${b.reissue||''}</td><td class="n">${b.warned||''}</td></tr>`).join('')}</table></div>`:'<div class="empty">이 기간에 기록이 없습니다.</div>';
  } else if(BT==='qual'){
    const P=D.pre;
    h=`<div class="qsum">${T.count?`${T.count}건 중 입력 오류 재발행 <b>${T.reissue_err}건</b>(${Math.round(100*T.reissue_err/T.count)}%) · 경고로 사전에 잡은 입력 <b>${T.warned}건</b>`:'이 기간 기록이 없습니다.'}</div>
      <div class="qgrid">
        <div class="qbox"><h3>수정 재발행 ${T.reissue}건</h3><table>${D.reasons.map(x=>`<tr><td>${x}${x==='입력 오류'?' <span class="why err">품질 이슈</span>':''}</td><td class="n">${T.by_reason[x]||0}</td></tr>`).join('')}</table></div>
        <div class="qbox"><h3>입력 이상치 경고 ${T.warned}건</h3>${Object.keys(T.warn_types).length?`<table>${Object.entries(T.warn_types).sort((a,b)=>b[1]-a[1]).map(([k,v])=>`<tr><td>${esc(k)}</td><td class="n">${v}</td></tr>`).join('')}</table>`:'<div class="hint">경고가 뜬 건이 없습니다.</div>'}<div class="hint" style="margin-top:6px">경고가 떴지만 확인하고 그대로 발행한 건 ${T.warned_acked}건</div></div>
        <div class="qbox"><h3>도입 전 비교</h3>${P?`<table><tr><td>도입 전 (${esc(P.from)} ~ ${esc(P.to)})</td><td class="n">${P.reissue}/${P.total}건 · <b>${P.pct}%</b></td></tr><tr><td>도입 후 (이 기간)</td><td class="n">${T.reissue_err}/${T.count||0}건 · <b>${T.count?Math.round(100*T.reissue_err/T.count):0}%</b></td></tr></table><div class="hint" style="margin-top:6px">도입 전 ${P.pct}% → 도입 후 ${T.count?Math.round(100*T.reissue_err/T.count):0}% · 출처: ${esc(P.source)||'미기재'}</div>`:'<div class="hint">도입 전 재발행·오류 건수가 입력되지 않았습니다. <a href="/실적/설정">기준 시간 설정 → 도입 전 비교</a>에 적으면 여기 비교가 나옵니다.</div>'}</div>
      </div>`;
  } else {
    const b=D.bases;
    h=`<table><tr><th>종류</th><th class="n">기준 시간(분, 시료 1종)</th><th class="n">시료 1종 추가마다</th><th>상태</th><th>근거</th></tr>
      ${D.kinds.map(([k,n])=>`<tr><td>${n}</td><td class="n"><b>${b[k].min}</b></td><td class="n">${S.per_sample[k]||'—'}</td><td>${b[k].status==='확정'?'<span class="why ok">확정</span>':`<span class="why">${b[k].status}</span>`}</td><td>${b[k].n?`실측 ${b[k].n}건 ${b[k].method} · ${esc(b[k].by.join(', '))} · 마지막 ${esc(b[k].last)}`:'실측 없음 — 손으로 넣은 값'}</td></tr>`).join('')}</table>
      <div class="rules" style="margin-top:14px">
      · <b>계산식</b> 절감 시간 = Σ( 기준 시간 + 시료 보정 × (시료 수 − 1) ) − 프로그램 입력 시간(${S.program}분) × 건수. 기록마다 그 시점의 기준 시간이 함께 저장되어 기준이 바뀌어도 과거 집계는 당시 값으로 유지됩니다.<br>
      · <b>발행 건수</b>는 성적서번호(${S.prefix}-YY-NNN)가 적힌 성적서를 번호별로 한 건씩 셉니다. 같은 번호를 다시 만들면 건수는 그대로이고 «수정 재발행»으로 따로 셉니다.<br>
      · 번호가 없는 것, test·테스트·임시·연습·샘플이 든 번호, 「연습용」, 담당자가 제외한 것은 세지 않습니다. 소급 등록은 발급일자 기준입니다.<br>
      · <b>전망</b>은 최근 3개월 평균 절감 시간 × 남은 개월입니다. 자료가 3개월 미만이면 보이지 않습니다.</div>
      <h3 style="font-size:12px;color:var(--mute);margin:16px 0 6px;letter-spacing:.06em">기준 시간 변경 이력</h3>
      ${S.history.length?`<table><tr><th>일시</th><th>누가</th><th>종류</th><th class="n">전</th><th class="n">후</th><th>방법</th></tr>${S.history.map(x=>`<tr><td>${esc(x.ts)}</td><td>${esc(x.by)}</td><td>${esc((D.kinds.find(([k])=>k===x.kind)||[,x.kind])[1])}</td><td class="n">${x.old}</td><td class="n">${x.new}</td><td>${esc(x.how)} · ${esc(x.status||'')}</td></tr>`).join('')}</table>`:'<div class="hint">아직 변경 이력이 없습니다.</div>'}`;
  }
  $('#bview').innerHTML=h;
}
document.addEventListener('mousemove',e=>{ const t=e.target.closest('[data-tip]'); const tip=$('#tip'); if(!t||t.classList.contains('kpi')){tip.style.display='none';return;} tip.textContent=t.dataset.tip; tip.style.display='block'; tip.style.left=(e.clientX+14)+'px'; tip.style.top=(e.clientY+14)+'px'; });
document.addEventListener('click',async e=>{
  const bd=e.target.closest('.badge'); if(bd){ OPENB=OPENB===bd.dataset.b?'':bd.dataset.b; renderBadgeDetail(); return; }
  const ct=e.target.closest('#ctabs button'); if(ct){ CT=ct.dataset.c; renderChart(); return; }
  const bt=e.target.closest('#btabs button'); if(bt){ BT=bt.dataset.b; renderBottom(); return; }
  const g=e.target.closest('#group button'); if(g){ G=g.dataset.g; $$('#group button').forEach(x=>x.classList.toggle('on',x===g)); load(); return; }
  const q=e.target.closest('#quick button'); if(q){ setQuick(q.dataset.q); load(); return; }
  if(e.target.id==='present'){ document.body.classList.toggle('present'); return; }
  if(e.target.id==='print'){ window.print(); return; }
  const x=e.target.closest('button.x'); if(x){ await fetch('/실적/제외',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({id:+x.dataset.id,excluded:x.dataset.ex==='1'})}); load(); }
});
document.addEventListener('change',async e=>{ const s=e.target.closest('select.rs'); if(s){ await fetch('/실적/사유',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({id:+s.dataset.id,reason:s.value})}); load(); } });
document.addEventListener('keydown',e=>{ if(e.key==='Escape') document.body.classList.remove('present'); });
['#from','#to'].forEach(s=>$(s).addEventListener('change',()=>{$$('#quick button').forEach(b=>b.classList.remove('on')); load();}));
if(new URLSearchParams(location.search).get('present')==='1') document.body.classList.add('present');
setQuick('year'); load();
</script></body></html>"""


SETTINGS_PAGE = r"""<!doctype html>
<html lang="ko"><head><meta charset="utf-8"><title>기준 시간 설정 · 실적</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
<link rel="stylesheet" href="https://cdn.jsdelivr.net/gh/orioncactus/pretendard@v1.3.9/dist/web/variable/pretendardvariable-dynamic-subset.min.css" onerror="this.remove()">
<style>
:root{--ink:#0f0f11;--ink-2:#3a3a40;--mute:#84848a;--faint:#b4b4b9;--line:#e3e0d9;--line-2:#eeebe4;--paper:#fff;--bg:#f1efea;--black:#0a0a0b;--ok:#1f5a3a;--ok-soft:#e6efe9}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);font:13px/1.5 "Pretendard Variable",Pretendard,"Malgun Gothic",system-ui,sans-serif}
__NAVCSS__
.wrap{max-width:960px;margin:0 auto;padding:36px 40px 80px}
h1{font-size:22px;font-weight:700;margin:0 0 4px}
.sub{color:var(--mute);font-size:12px;margin-bottom:24px;line-height:1.6}
.card{background:var(--paper);border:1px solid var(--line);padding:22px 24px;margin-bottom:18px}
.card h2{font-size:14px;font-weight:700;margin:0 0 4px}
.card .hint{font-size:11.5px;color:var(--mute);margin-bottom:14px;line-height:1.55}
.kind{border-top:1px solid var(--line);padding:14px 0 6px}
.kind .kh{display:flex;justify-content:space-between;align-items:baseline;flex-wrap:wrap;gap:8px;margin-bottom:8px}
.kind .kh b{font-size:13.5px}
.kind .kh .st{font-size:12px;color:var(--mute)}
.kind .kh .st .ok{color:var(--ok);font-weight:600}
.kind .kh .st .est{color:var(--ink-2);font-weight:600}
table{width:100%;border-collapse:collapse;font-size:12.5px}
th{font-size:10.5px;letter-spacing:.1em;color:var(--mute);font-weight:600;text-align:left;padding:5px 6px;border-bottom:1px solid var(--line)}
td{padding:4px 6px;border-bottom:1px solid var(--line-2)}
td input{font:inherit;width:100%;padding:6px 8px;border:1px solid var(--line);background:#fff}
td button{font:inherit;font-size:11px;padding:4px 8px;border:1px solid var(--line);background:#fff;cursor:pointer;color:var(--ink-2)}
.add{font:inherit;font-size:12px;padding:6px 10px;border:1px solid var(--line);background:#fff;cursor:pointer;margin-top:6px}
.two{display:grid;grid-template-columns:1fr 1fr;gap:14px}
.three{display:grid;grid-template-columns:1fr 1fr 1fr;gap:14px}
label{display:block;font-size:10.5px;letter-spacing:.1em;color:var(--mute);font-weight:600;margin-bottom:5px}
.f input,.f select{font:inherit;width:100%;padding:8px;border:1px solid var(--line);background:#fff}
.save{font:inherit;font-size:12.5px;padding:10px 20px;background:var(--black);color:#fff;border:none;cursor:pointer}
.msg{font-size:12px;color:var(--ok);margin-left:12px}
.rowline{display:flex;gap:14px;align-items:flex-end;flex-wrap:wrap;margin-top:12px}
.rowline .f{min-width:180px}
.vtab td,.vtab th{font-size:12px}
.ok{color:var(--ok)}.bad{color:#3a3a40;font-weight:600}
@media(max-width:800px){.two,.three{grid-template-columns:1fr}.wrap{padding:24px 16px 60px}}
</style></head><body>
__NAV__
<div class="wrap">
<h1>기준 시간 설정</h1>
<div class="sub">실적의 절감 시간은 «손으로 만들 때 걸리던 시간 − 프로그램 입력 시간»입니다. 손으로 만들 때 시간은 담당자가 실제로 잰 기록의 평균(또는 중앙값)으로 정합니다. 실측이 종류별 2건 이상이면 «확정», 1건이면 «확인 중», 없으면 손으로 넣은 «추정치»를 씁니다. 기준이 바뀌어도 과거 집계는 당시 값으로 유지됩니다.</div>

<div class="card"><h2>실측 기록 — 성적서 한 장을 손으로 만드는 데 실제 걸린 시간</h2>
  <div class="hint">담당자 이름 · 측정일 · 걸린 시간(분) · 그때 시료 수 · 메모. 시료 수가 1종이 아니면 «시료 1종 추가마다» 값으로 1종 기준으로 환산해 평균을 냅니다.</div>
  <div class="rowline"><div class="f"><label>기준 시간 계산 방법</label><select id="method"><option value="mean">평균</option><option value="median">중앙값</option></select></div>
    <div class="f"><label>프로그램에 입력하는 시간 (분)</label><input id="program" inputmode="decimal"></div>
    <div class="f"><label>이 설정을 바꾸는 사람 (이력에 남음)</label><input id="who" placeholder="이름"></div></div>
  <div id="kinds"></div></div>

<div class="card"><h2>도입 전 비교 (관리자 입력)</h2>
  <div class="hint">수기 작성 시절의 재발행·오류 건수를 적으면 품질 기록 탭에 «도입 전 N% → 도입 후 N%» 비교가 나옵니다. 비워 두면 그 블록은 숨깁니다.</div>
  <div class="three"><div class="f"><label>기간 시작</label><input type="date" id="pre_from"></div><div class="f"><label>기간 끝</label><input type="date" id="pre_to"></div><div></div>
    <div class="f"><label>발행 성적서 수</label><input id="pre_total" inputmode="numeric"></div><div class="f"><label>그중 재발행·오류 건수</label><input id="pre_reissue" inputmode="numeric"></div><div class="f"><label>출처 메모</label><input id="pre_source" placeholder="예: 2025년 성적서 대장 (품질팀)"></div></div></div>

<div class="card"><h2>그 밖의 기준</h2>
  <div class="three"><div class="f"><label>하루 근무시간 (근무일 환산)</label><input id="workday_hours" inputmode="decimal"></div>
    <div class="f"><label>성적서번호 접두어</label><input id="prefix" maxlength="4"></div><div class="f"><label>다음 번호 미리보기</label><input id="nextno" readonly style="background:#f7f5f0"></div>
    <div class="f"><label>도입 전 올해 손으로 만든 건수 (참고)</label><input id="baseline" inputmode="numeric"></div><div class="f" style="grid-column:span 2"><label>메모</label><input id="baseline_note" placeholder="예: L-26-042까지 손으로 작성"></div></div></div>

<button class="save" id="save">저장</button><span class="msg" id="msg"></span>

<div class="card" id="verify" style="margin-top:28px"><h2>계산 검증</h2><div class="hint" id="vhint"></div><div id="vbody"></div></div>
</div>
<script>
const $=(s,r=document)=>r.querySelector(s), $$=(s,r=document)=>[...r.querySelectorAll(s)];
const KINDS=__KINDS__; let S=null, D=null;
function esc(s){return String(s??'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));}
async function load(){
  D=await (await fetch('/실적/data?from=2000-01-01&to=2000-01-02&group=year')).json(); S=await (await fetch('/실적/설정값')).json();
  $('#method').value=S.method||'mean'; $('#program').value=S.program;
  $('#kinds').innerHTML=KINDS.map(([k,n,per])=>{ const b=D.bases[k]; const ms=S.measurements[k]||[];
    return `<div class="kind" data-k="${k}"><div class="kh"><b>${n}</b><span class="st">기준 시간 <b>${b.min}분</b> · ${b.status==='확정'?'<span class="ok">확정</span>':b.status==='확인 중'?'<span class="est">확인 중 (실측 1건 — 1건 더 있으면 확정)</span>':'<span class="est">추정치 (실측 없음)</span>'} · 실측 ${b.n}건</span></div>
      <table><tr><th>담당자</th><th>측정일</th><th>걸린 시간(분)</th><th>시료 수</th><th>메모</th><th></th></tr><tbody class="ms">${ms.map(m=>row(m)).join('')}</tbody></table>
      <button class="add">+ 실측 기록 추가</button>
      <div class="rowline"><div class="f"><label>실측이 없을 때 쓰는 손 입력값 (분, 시료 1종) — 추정치</label><input class="manual" value="${S.manual[k]}" inputmode="decimal"></div>
        ${per?`<div class="f"><label>시료 1종 추가마다 (분)</label><input class="per" value="${S.per_sample[k]}" inputmode="decimal"></div>`:'<div class="f"><label>시료 1종 추가마다</label><input value="— (한 장에 시료 하나)" disabled></div>'}</div></div>`; }).join('');
  const P=S.pre||{}; $('#pre_from').value=P.from||''; $('#pre_to').value=P.to||''; $('#pre_total').value=P.total||''; $('#pre_reissue').value=P.reissue||''; $('#pre_source').value=P.source||'';
  $('#workday_hours').value=S.workday_hours||8; $('#prefix').value=S.prefix||'L'; $('#baseline').value=S.baseline||0; $('#baseline_note').value=S.baseline_note||'';
  try{ const nn=await (await fetch('/성적서/번호')).json(); $('#nextno').value=nn.suggest; }catch(e){}
  loadVerify();
}
function row(m){ m=m||{}; return `<tr><td><input class="by" value="${esc(m.by||'')}"></td><td><input type="date" class="date" value="${esc(m.date||'')}"></td><td><input class="minutes" inputmode="decimal" value="${m.minutes||''}"></td><td><input class="ns" inputmode="numeric" value="${m.n_samples||1}"></td><td><input class="memo" value="${esc(m.memo||'')}"></td><td><button class="rm">빼기</button></td></tr>`; }
document.addEventListener('click',async e=>{
  if(e.target.classList.contains('add')){ $('.ms',e.target.closest('.kind')).insertAdjacentHTML('beforeend',row()); return; }
  if(e.target.classList.contains('rm')){ e.target.closest('tr').remove(); return; }
  if(e.target.id==='save'){
    const p={method:$('#method').value,program:$('#program').value,by:$('#who').value,manual:{},per_sample:{},measurements:{},
      pre:{from:$('#pre_from').value,to:$('#pre_to').value,total:$('#pre_total').value,reissue:$('#pre_reissue').value,source:$('#pre_source').value},
      workday_hours:$('#workday_hours').value,prefix:$('#prefix').value,baseline:$('#baseline').value,baseline_note:$('#baseline_note').value};
    $$('.kind').forEach(kd=>{ const k=kd.dataset.k; p.manual[k]=$('.manual',kd).value; const per=$('.per',kd); if(per) p.per_sample[k]=per.value;
      p.measurements[k]=$$('.ms tr',kd).map(tr=>({by:$('.by',tr).value,date:$('.date',tr).value,minutes:$('.minutes',tr).value,n_samples:$('.ns',tr).value,memo:$('.memo',tr).value})).filter(m=>m.minutes); });
    const r=await fetch('/실적/설정',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(p)});
    $('#msg').textContent=r.ok?'저장했습니다. 실적 화면에 바로 반영됩니다.':'저장하지 못했습니다.'; setTimeout(()=>$('#msg').textContent='',4000); load(); return; }
  if(e.target.id==='verifyRun'){ e.target.disabled=true; await fetch('/실적/검증',{method:'POST'}); setTimeout(loadVerify,3000); setTimeout(loadVerify,9000); }
});
async function loadVerify(){
  let V; try{ V=await (await fetch('/실적/검증')).json(); }catch(e){ return; }
  const K={finished:'라돈·토론 완제품',material:'라돈 원자재',kc:'매트리스 KC',ks:'매트리스 KS'};
  $('#vhint').innerHTML=`과거에 사람이 발행한 성적서에 적힌 값과, 같은 원본으로 프로그램이 계산한 값을 항목별로 대조합니다. ${V.at?`<b>${V.ok} / ${V.n} 일치</b> · ${esc(V.at)}${V.stale?' · 계산 코드가 그 뒤 바뀌어 다시 점검이 필요합니다':''}`:'아직 점검하지 않았습니다.'} <button class="add" id="verifyRun" ${V.busy?'disabled':''}>${V.busy?'점검 중…':'지금 다시 점검'}</button>`;
  $('#vbody').innerHTML=(V.groups||[]).map(g=>`<div style="font-weight:600;margin:12px 0 4px">${K[g.kind]||g.kind} <span style="font-weight:400;color:var(--mute);font-size:12px">${g.rows.length?`${g.rows.filter(r=>r.ok).length}/${g.rows.length} 일치`:''}${g.note?` · ${esc(g.note)}`:''}${g.error?` · ${esc(g.error)}`:''}</span></div>${g.rows.length?`<table class="vtab"><tr><th>항목</th><th>사람값</th><th>프로그램값</th><th>일치</th></tr>${g.rows.map(r=>`<tr><td>${esc(r.item)}</td><td>${esc(JSON.stringify(r.human))}</td><td>${esc(JSON.stringify(r.program))}</td><td>${r.ok?'<span class="ok">일치</span>':'<span class="bad">다름</span>'}</td></tr>`).join('')}</table>`:''}`).join('');
}
load();
</script></body></html>""".replace("__KINDS__", json.dumps([(k, KIND_NAMES[k], k in ("finished", "material")) for k in KIND_ORDER], ensure_ascii=False))


def page(build=""):
    import portal
    return PAGE.replace("__NAV__", portal.nav("/실적", '<a href="/실적/설정">기준 시간 설정</a>', build)).replace("__NAVCSS__", portal.NAV_CSS)


def settings_page(build=""):
    import portal
    return SETTINGS_PAGE.replace("__NAV__", portal.nav("/실적", '<a href="/실적">‹ 실적으로</a>', build)).replace("__NAVCSS__", portal.NAV_CSS)


try:
    migrate()
except Exception:
    pass
